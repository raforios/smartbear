'''
    Converts a Markdown deliverable into a Word (.docx) document.

    Written for the client-facing reports (TRADE / BINARIA, informes del
    Ministerio): those are drafted in Markdown inside the repo and sent as
    Word, so this keeps both in sync instead of re-typing the document.

    Supports the subset actually used in those documents: ATX headings,
    paragraphs with inline bold / italic / code, bullet and numbered lists,
    GitHub-flavoured tables, fenced code blocks and horizontal rules.

    Usage:
        python tools/md_to_docx.py <input.md> [output.docx]
'''
import argparse
import os
import re
from typing import List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CODE_FONT = 'Consolas'
CODE_SHADING = 'F2F3F5'
HEADER_SHADING = 'E3ECF8'
# Bold / italic / inline code, captured so the text keeps its emphasis.
_INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)')
# A line opening with a bold label ("**Cliente:** BINARIA") is a field of its
# own, not the continuation of the line above it.
_LABEL = re.compile(r'^\*\*[^*]+:\*\*')


def _shade(cell_or_paragraph, color: str) -> None:
    '''
        Paints a solid background on a table cell or a paragraph.

        Args:
            cell_or_paragraph: python-docx cell or paragraph to shade.
            color (str): Hex color without '#'.
    '''
    element = OxmlElement('w:shd')
    element.set(qn('w:val'), 'clear')
    element.set(qn('w:fill'), color)
    target = getattr(cell_or_paragraph, '_tc', None)
    if target is not None:
        target.get_or_add_tcPr().append(element)
        return
    cell_or_paragraph.paragraph_format.element.get_or_add_pPr().append(element)


def _add_rich_text(paragraph, text: str, bold: bool = False) -> None:
    '''
        Writes text into a paragraph, honouring inline **bold**, *italic* and
        `code` spans.

        Args:
            paragraph: Destination paragraph.
            text (str): Markdown inline text.
            bold (bool): Force the whole run bold (table headers).
    '''
    for chunk in _INLINE.split(text):
        if not chunk:
            continue
        run = paragraph.add_run()
        if chunk.startswith('**') and chunk.endswith('**'):
            run.text = chunk[2:-2]
            run.bold = True
        elif chunk.startswith('`') and chunk.endswith('`'):
            run.text = chunk[1:-1]
            run.font.name = CODE_FONT
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        elif chunk.startswith('*') and chunk.endswith('*'):
            run.text = chunk[1:-1]
            run.italic = True
        else:
            run.text = chunk
        if bold:
            run.bold = True


def _add_code_block(document: Document, lines: List[str]) -> None:
    '''
        Renders a fenced code block as a shaded monospace paragraph.
    '''
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run('\n'.join(lines))
    run.font.name = CODE_FONT
    run.font.size = Pt(8.5)
    _shade(paragraph, CODE_SHADING)


def _split_row(line: str) -> List[str]:
    '''
        Splits a Markdown table row into its trimmed cells.
    '''
    return [cell.strip() for cell in line.strip().strip('|').split('|')]


def _add_table(document: Document, rows: List[str]) -> None:
    '''
        Renders a GitHub-flavoured table, shading the header row.

        Args:
            document (Document): Target document.
            rows (list[str]): Raw markdown lines, separator row included.
    '''
    header = _split_row(rows[0])
    body = [_split_row(row) for row in rows[2:]]
    table = document.add_table(rows = 1, cols = len(header))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, title in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ''
        _add_rich_text(cell.paragraphs[0], title, bold = True)
        _shade(cell, HEADER_SHADING)

    for line in body:
        cells = table.add_row().cells
        for index, value in enumerate(line[:len(header)]):
            cells[index].text = ''
            _add_rich_text(cells[index].paragraphs[0], value)
    document.add_paragraph()


def _add_wrapped(document: Document, buffer: List[str]) -> None:
    '''
        Emits a soft-wrapped Markdown paragraph as ONE Word paragraph.

        Markdown treats consecutive lines as the same paragraph; without this
        every line break in the source became its own paragraph in Word and the
        document read as a broken column of fragments.
    '''
    kind, text = buffer[0], ' '.join(buffer[1:])
    if kind == 'bullet':
        paragraph = document.add_paragraph(style = 'List Bullet')
    elif kind == 'number':
        paragraph = document.add_paragraph(style = 'List Number')
    else:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_rich_text(paragraph, text)


def _flush(document: Document, buffer: List[str], kind: str) -> None:
    '''
        Emits whatever block the parser had accumulated.
    '''
    if not buffer:
        return
    if kind == 'table':
        _add_table(document, buffer)
    elif kind == 'code':
        _add_code_block(document, buffer)
    elif kind == 'text':
        _add_wrapped(document, buffer)
    buffer.clear()


class _Renderer:
    '''
        Line-by-line Markdown reader that writes into a Word document.

        Holds the open block (code fence, table or wrapped paragraph) so the
        line handlers stay small: each one decides whether the current block
        ends and what to open next.
    '''

    def __init__(self, document: Document):
        self.document = document
        self.buffer: List[str] = []
        self.block: Optional[str] = None

    def close(self) -> None:
        '''
            Emits the block still open, if any.
        '''
        _flush(self.document, self.buffer, self.block or '')
        self.block = None

    def feed(self, line: str) -> None:
        '''
            Consumes one source line.

            Args:
                line (str): Raw line, indentation included (code needs it).
        '''
        stripped = line.strip()
        if self._handle_fence(line, stripped):
            return
        if self._handle_table(stripped):
            return
        if self._handle_break(stripped):
            return
        if self._handle_heading(stripped):
            return
        if self._handle_list(stripped):
            return
        self._handle_paragraph(stripped)

    def _handle_fence(self, line: str, stripped: str) -> bool:
        '''
            Opens or closes a fenced code block; inside one, everything is
            captured verbatim.
        '''
        if stripped.startswith('```'):
            opening = self.block != 'code'
            self.close()
            self.block = 'code' if opening else None
            return True
        if self.block == 'code':
            self.buffer.append(line)
            return True
        return False

    def _handle_table(self, stripped: str) -> bool:
        '''
            Accumulates the rows of a GitHub-flavoured table.
        '''
        if stripped.startswith('|') and stripped.endswith('|'):
            if self.block != 'table':
                self.close()
                self.block = 'table'
            self.buffer.append(stripped)
            return True
        if self.block == 'table':
            self.close()
        return False

    def _handle_break(self, stripped: str) -> bool:
        '''
            Blank lines and horizontal rules end the current paragraph; an
            explicit `&nbsp;` line is a deliberate blank line (signature
            blocks) and must survive into the document.
        '''
        if not stripped:
            self.close()
            return True
        if stripped == '&nbsp;':
            self.close()
            self.document.add_paragraph()
            return True
        if stripped.startswith('---') and set(stripped) == {'-'}:
            self.close()
            self.document.add_paragraph()
            return True
        return False

    def _handle_heading(self, stripped: str) -> bool:
        '''
            Writes an ATX heading (levels 1 to 4).
        '''
        heading = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if not heading:
            return False
        self.close()
        paragraph = self.document.add_heading('', level = len(heading.group(1)))
        _add_rich_text(paragraph, heading.group(2))
        return True

    def _handle_list(self, stripped: str) -> bool:
        '''
            Opens a bullet or numbered item; further lines wrap into it.
        '''
        bullet = re.match(r'^[-*]\s+(.*)$', stripped)
        numbered = re.match(r'^\d+\.\s+(.*)$', stripped)
        if not (bullet or numbered):
            return False
        self.close()
        self.block = 'text'
        self.buffer.extend([
            'bullet' if bullet else 'number',
            (bullet or numbered).group(1),
        ])
        return True

    def _handle_paragraph(self, stripped: str) -> None:
        '''
            Starts a paragraph or wraps the line into the open one.
        '''
        if self.block == 'text' and _LABEL.match(stripped):
            self.close()
        if self.block != 'text':
            self.close()
            self.block = 'text'
            self.buffer.append('paragraph')
        self.buffer.append(stripped)


def convert(markdown_path: str, docx_path: str) -> str:
    '''
        Converts a Markdown file into a .docx document.

        Args:
            markdown_path (str): Source .md file.
            docx_path (str): Destination .docx file.

        Returns:
            str: The path written.
    '''
    document = Document()
    style = document.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    renderer = _Renderer(document)
    with open(markdown_path, encoding = 'utf-8') as handle:
        for line in handle.read().splitlines():
            renderer.feed(line)
    renderer.close()

    document.save(docx_path)
    return docx_path


def main() -> None:
    '''
        CLI entry point.
    '''
    parser = argparse.ArgumentParser(description = __doc__)
    parser.add_argument('source', help = 'Markdown file to convert.')
    parser.add_argument('target', nargs = '?', help = 'Output .docx (defaults alongside).')
    args = parser.parse_args()

    target = args.target or f'{os.path.splitext(args.source)[0]}.docx'
    print(f'-- {convert(args.source, target)}')


if __name__ == '__main__':
    main()
